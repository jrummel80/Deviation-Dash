from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

try:
    import python_calamine  # noqa: F401

    HAS_CALAMINE = True
except ImportError:
    HAS_CALAMINE = False

MONTHLY_COUNT = 24
MONTH_LABEL_PATTERN = re.compile(r"^[A-Za-z]{3}-\d{2}$")
WORKBOOK_CACHE_VERSION = "2026-03-23-calamine-v1"
WORKBOOK_CACHE_DIR = Path(__file__).resolve().parent.parent / ".cache" / "workbooks"

HEADER_ALIASES = {
    "supplier": ["Supplier"],
    "prod_group": ["Prod Group"],
    "price_group": ["Price Group"],
    "item": ["Item"],
    "description": ["Description"],
    "location": ["Location"],
    "frequency": ["Frequency"],
    "lead_time": ["Lead Time"],
    "uom": ["UOM"],
    "uom_size": ["UOM Size"],
    "type": ["Type"],
    "buyable": ["Buyable"],
    "sub_flag": ["Sub?"],
    "stockable": ["Stockable"],
    "br": ["BR", "BR Flag"],
    "job_usage": ["Job Usage"],
    "override": ["Override"],
    "per": ["Per"],
    "override_date": ["Ovr Date", "Oride Date"],
    "stock_room": ["Stk Rm"],
    "show_room": ["ShwRm"],
    "status": ["Status"],
    "volume": ["Volume"],
    "abc": ["ABC"],
    "season": ["Season"],
    "sensitivity": ["Sensitivity"],
    "date_created": ["Date Created"],
    "last_receipt": ["Last Receipt"],
    "last_sale": ["Last Sale"],
    "mac": ["MAC"],
    "line_cost": ["Line Cost"],
    "gross_qoh": ["Gross QOH"],
    "net_qoh": ["Net QOH"],
    "inbound": ["Inbound"],
    "min_value": ["Min"],
    "current_max": ["Max"],
    "supplier_min_amount": ["S. Min Amt", "S Min Amt"],
    "package_qty": ["Pkg Qty"],
    "locations_with_max": ["Locs with Max"],
    "ttmd": ["TTMD2", "TTMD"],
    "l6m_summer": ["L6M Sumr"],
    "l6m_winter": ["L6M Wint"],
    "usage_24m_total": ["1-24M"],
    "usage_12m_total": ["1-12M"],
    "usage_13_24m_total": ["13-24M"],
    "legacy_mean": ["Mean"],
}

NUMERIC_COLUMNS = {
    "lead_time",
    "mac",
    "line_cost",
    "gross_qoh",
    "net_qoh",
    "inbound",
    "min_value",
    "current_max",
    "supplier_min_amount",
    "package_qty",
    "locations_with_max",
    "ttmd",
    "l6m_summer",
    "l6m_winter",
    "usage_24m_total",
    "usage_12m_total",
    "usage_13_24m_total",
    "legacy_mean",
}

TEXT_COLUMNS = {
    "supplier",
    "prod_group",
    "price_group",
    "item",
    "description",
    "location",
    "frequency",
    "uom",
    "uom_size",
    "type",
    "buyable",
    "sub_flag",
    "stockable",
    "br",
    "job_usage",
    "override",
    "per",
    "stock_room",
    "show_room",
    "status",
    "volume",
    "abc",
    "season",
    "sensitivity",
}


@dataclass(frozen=True)
class LoadedWorkbook:
    data: pd.DataFrame
    monthly_columns: list[str]
    source_name: str


def load_data_workbook(file_bytes: bytes, source_name: str) -> LoadedWorkbook:
    cache_key = _workbook_cache_key(file_bytes)
    cached = _load_cached_workbook(cache_key, source_name)
    if cached is not None:
        return cached

    if HAS_CALAMINE:
        try:
            extracted_frame, monthly_columns = _load_with_calamine(file_bytes)
            normalized = _normalize_data_frame(extracted_frame, monthly_columns)
            loaded = LoadedWorkbook(data=normalized, monthly_columns=monthly_columns, source_name=source_name)
            _save_cached_workbook(cache_key, loaded)
            return loaded
        except Exception:
            pass

    normalized, monthly_columns = _load_with_openpyxl(file_bytes)
    loaded = LoadedWorkbook(data=normalized, monthly_columns=monthly_columns, source_name=source_name)
    _save_cached_workbook(cache_key, loaded)
    return loaded


def _load_with_calamine(file_bytes: bytes) -> tuple[pd.DataFrame, list[str]]:
    raw_frame = pd.read_excel(BytesIO(file_bytes), sheet_name="Data", engine="calamine")
    if raw_frame.empty and len(raw_frame.columns) == 0:
        raise ValueError("The 'Data' sheet is empty.")
    return _extract_relevant_frame(raw_frame, tuple(raw_frame.columns))


def _load_with_openpyxl(file_bytes: bytes) -> tuple[pd.DataFrame, list[str]]:
    workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True, read_only=True)

    if "Data" not in workbook.sheetnames:
        raise ValueError("Workbook must include a 'Data' sheet.")

    sheet = workbook["Data"]
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)

    if header_row is None:
        raise ValueError("The 'Data' sheet is empty.")

    header_lookup = _header_lookup(header_row)
    mapped_indexes = {
        canonical: _first_matching_index(header_lookup, aliases)
        for canonical, aliases in HEADER_ALIASES.items()
    }
    monthly_columns, monthly_indexes = _extract_month_columns(header_row, header_lookup)

    missing = [column for column in ("supplier", "item", "location") if mapped_indexes[column] is None]
    if missing:
        raise ValueError(f"Workbook is missing required columns: {', '.join(missing)}")
    if not monthly_columns:
        raise ValueError("Workbook is missing monthly demand columns.")

    records: list[dict[str, object]] = []
    for row in rows:
        row_values = list(row)
        if not any(value not in (None, "") for value in row_values):
            continue

        record = {
            canonical: _value_at(row_values, column_index)
            for canonical, column_index in mapped_indexes.items()
        }
        for month_label, column_index in zip(monthly_columns, monthly_indexes):
            record[month_label] = _value_at(row_values, column_index)
        records.append(record)

    data_frame = pd.DataFrame(records)
    normalized = _normalize_data_frame(data_frame, monthly_columns)
    return normalized, monthly_columns


def _extract_relevant_frame(
    data_frame: pd.DataFrame,
    header_row: tuple[object, ...],
) -> tuple[pd.DataFrame, list[str]]:
    header_lookup = _header_lookup(header_row)
    mapped_indexes = {
        canonical: _first_matching_index(header_lookup, aliases)
        for canonical, aliases in HEADER_ALIASES.items()
    }
    monthly_columns, monthly_indexes = _extract_month_columns(header_row, header_lookup)

    missing = [column for column in ("supplier", "item", "location") if mapped_indexes[column] is None]
    if missing:
        raise ValueError(f"Workbook is missing required columns: {', '.join(missing)}")
    if not monthly_columns:
        raise ValueError("Workbook is missing monthly demand columns.")

    extracted: dict[str, object] = {}
    for canonical, column_index in mapped_indexes.items():
        if column_index is None:
            extracted[canonical] = None
        else:
            extracted[canonical] = data_frame.iloc[:, column_index]
    for month_label, column_index in zip(monthly_columns, monthly_indexes):
        extracted[month_label] = data_frame.iloc[:, column_index]

    trimmed = pd.DataFrame(extracted).dropna(how="all").reset_index(drop=True)
    return trimmed, monthly_columns


def _workbook_cache_key(file_bytes: bytes) -> str:
    digest = hashlib.sha256(file_bytes).hexdigest()
    return f"{WORKBOOK_CACHE_VERSION}-{digest}"


def _cache_paths(cache_key: str) -> tuple[Path, Path]:
    return (
        WORKBOOK_CACHE_DIR / f"{cache_key}.parquet",
        WORKBOOK_CACHE_DIR / f"{cache_key}.json",
    )


def _load_cached_workbook(cache_key: str, source_name: str) -> LoadedWorkbook | None:
    parquet_path, metadata_path = _cache_paths(cache_key)
    if not parquet_path.exists() or not metadata_path.exists():
        return None

    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        monthly_columns = list(metadata.get("monthly_columns", []))
        data = pd.read_parquet(parquet_path)
        return LoadedWorkbook(data=data, monthly_columns=monthly_columns, source_name=source_name)
    except Exception:
        parquet_path.unlink(missing_ok=True)
        metadata_path.unlink(missing_ok=True)
        return None


def _save_cached_workbook(cache_key: str, loaded: LoadedWorkbook) -> None:
    WORKBOOK_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    parquet_path, metadata_path = _cache_paths(cache_key)
    loaded.data.to_parquet(parquet_path, index=False)
    metadata_path.write_text(
        json.dumps({"monthly_columns": loaded.monthly_columns}, indent=2),
        encoding="utf-8",
    )


def _header_lookup(header_row: tuple[object, ...]) -> dict[str, int]:
    lookup: dict[str, int] = {}
    for index, value in enumerate(header_row):
        key = _header_key(value)
        if key and key not in lookup:
            lookup[key] = index
    return lookup


def _first_matching_index(header_lookup: dict[str, int], aliases: list[str]) -> int | None:
    for alias in aliases:
        key = _header_key(alias)
        if key in header_lookup:
            return header_lookup[key]
    return None


def _extract_month_columns(
    header_row: tuple[object, ...],
    header_lookup: dict[str, int],
) -> tuple[list[str], list[int]]:
    raw_months: list[tuple[str, int]] = []

    for index, value in enumerate(header_row):
        label = str(value).strip() if value not in (None, "") else ""
        if MONTH_LABEL_PATTERN.match(label):
            raw_months.append((label, index))

    if not raw_months:
        usage_end_index = _first_matching_index(header_lookup, HEADER_ALIASES["usage_13_24m_total"])
        if usage_end_index is None:
            return [], []
        for offset in range(1, MONTHLY_COUNT + 1):
            index = usage_end_index + offset
            if index >= len(header_row):
                break
            label = str(header_row[index]).strip() if header_row[index] not in (None, "") else f"Month {offset:02d}"
            raw_months.append((label, index))

    labels = _dedupe_month_labels([label for label, _ in raw_months[:MONTHLY_COUNT]])
    indexes = [index for _, index in raw_months[:MONTHLY_COUNT]]
    return labels, indexes


def _dedupe_month_labels(raw_headers: list[str]) -> list[str]:
    cleaned: list[str] = []
    seen: set[str] = set()

    for index, value in enumerate(raw_headers, start=1):
        label = str(value).strip() if value not in (None, "") else f"Month {index:02d}"
        if label.startswith("="):
            label = f"Month {index:02d}"
        if label in seen:
            label = f"{label} ({index})"
        seen.add(label)
        cleaned.append(label)

    return cleaned


def _normalize_data_frame(data_frame: pd.DataFrame, monthly_columns: list[str]) -> pd.DataFrame:
    normalized = data_frame.copy()

    for canonical in HEADER_ALIASES:
        if canonical not in normalized.columns:
            normalized[canonical] = None

    for column in TEXT_COLUMNS:
        if column in normalized.columns:
            normalized[column] = normalized[column].fillna("").map(_clean_text)

    for column in NUMERIC_COLUMNS:
        if column in normalized.columns:
            normalized[column] = pd.to_numeric(normalized[column], errors="coerce").fillna(0.0)

    for column in monthly_columns:
        normalized[column] = pd.to_numeric(normalized[column], errors="coerce").fillna(0.0)

    for column in ("override_date", "date_created", "last_receipt", "last_sale"):
        if column in normalized.columns:
            normalized[column] = normalized[column].fillna("").astype(str).replace("NaT", "")

    normalized["location"] = normalized["location"].map(_clean_location)
    normalized["group_key"] = normalized["supplier"] + " | " + normalized["item"]
    return normalized


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text == "nan" else text


def _clean_location(value: object) -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _header_key(value: object) -> str:
    if value in (None, ""):
        return ""
    return "".join(character.lower() for character in str(value).strip() if character.isalnum())


def _value_at(row_values: list[object], index: int | None) -> object:
    if index is None or index >= len(row_values):
        return None
    return row_values[index]
